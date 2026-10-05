import streamlit as st
import pandas as pd
import io
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime
import re

st.set_page_config(page_title="Carga Rappi", layout="centered")
st.title("Carga de Datos Rappi")

uploaded_file = st.file_uploader("Seleccionar archivo Excel", type=['xlsx', 'xls'])

MAPEO_CODIGOS = {
    "Fridays Comas": "0000000016",
    "Fridays Mall del Sur": "0000000012",
    "Fridays Óvalo Gutiérrez": "0000000001",
    "Fridays Ovalo Gutierrez": "0000000001",
    "Fridays Ovalo Monitor": "0000000019",
    "Fridays Óvalo Monitor": "0000000019",
    "Fridays Plaza Lima Norte": "0000000017",
    "Fridays Primavera": "0000000009",
    "Fridays Rambla": "0000000010",
    "Fridays Salaverry": "0000000008",
    "Fridays San Miguel": "0000000003",
    "Fridays Santa Anita": "0000000015",
    "Fridays DK La Molina": "0000000018",
    "Fridays Dk La Molina": "0000000018",
    "Fridays Puruchuco": "0000000014",
    "Fridays San Juan de Lurigancho": "0000000020",
    "Fridays Piura": "0000000024",
    "Fridays - Real Plaza Arequipa": "0000000006",
    "Fridays - Mall Aventura Porongoche": "0000000023",
    "Fridays Mall Aventura Porongoche": "0000000023",
    "Fridays El Polo": "0000000027",
    "Fridays - Cusco": "0000000025",
    "Fridays Piura - Real Plaza": "0000000024",
    "Smash Burger by Fridays Óvalo Gutiérrez": "0000000001",
    "Smash Burger by Fridays Ovalo Gutierrez": "0000000001",
    "Smash Burger by Fridays Ovalo Monitor": "0000000019",
    "Smash Burger by Fridays Óvalo Monitor": "0000000019",
    "Smash Burger by Fridays San Juan de Lurigancho": "0000000020",
    "Smash Burger by Fridays Dk la Molina": "0000000018",
    "Smash Burger by Fridays DK La Molina": "0000000018",
    "Smash Burger by Fridays Santa Anita": "0000000015",
    "Smash Burger by Fridays Piura": "0000000024",
    "Smash Burger by Fridays Plaza Lima Norte": "0000000017",
    "Smash Burger by Fridays San Miguel": "0000000003",
    "Smash Burger By Fridays - Mall Aventura Porongoche": "0000000023",
    "Smash Burger by Fridays Mall Aventura Porongoche": "0000000023",
    "Smash Burger by Fridays Salaverry": "0000000008",
    "Smash Burger by Fridays Mall del Sur": "0000000012",
    "Smash Burger by Fridays Rambla": "0000000010",
    "Smash Burger by Fridays Comas": "0000000016",
    "Smash Burger by Fridays Puruchuco": "0000000014",
    "Smash Burger by Fridays Arequipa": "0000000006",
    "Smash Burger by Fridays Primavera": "0000000009",
    "Smash Burger by Fridays El Polo": "0000000027",
    "Smash Burger by Fridays Cusco": "0000000025",
}

def obtener_codigo_tienda(nombre_tienda):
    if pd.isna(nombre_tienda):
        return ""
    nombre_tienda = str(nombre_tienda).strip()
    if nombre_tienda in MAPEO_CODIGOS:
        return MAPEO_CODIGOS[nombre_tienda]
    nombre_normalizado = nombre_tienda.lower()
    for nombre_mapeado, codigo in MAPEO_CODIGOS.items():
        if nombre_mapeado.lower() == nombre_normalizado:
            return codigo
    return ""

def convertir_fecha_manual(valor):
    if pd.isna(valor):
        return None
    if isinstance(valor, (datetime, pd.Timestamp)):
        return pd.Timestamp(valor).normalize()
    texto = str(valor).strip()
    if "," in texto:
        texto = texto.split(",")[0].strip()
    match = re.search(r'(\d{1,2})\s+([a-záéíóúñ]+)\.?\s+(\d{4})', texto, re.IGNORECASE)
    if not match:
        return None
    dia = int(match.group(1))
    mes_str = match.group(2).lower()
    anio = int(match.group(3))
    meses = {
        'ene': 1, 'enero': 1, 'feb': 2, 'febrero': 2, 'mar': 3, 'marzo': 3,
        'abr': 4, 'abril': 4, 'may': 5, 'mayo': 5, 'jun': 6, 'junio': 6,
        'jul': 7, 'julio': 7, 'ago': 8, 'agosto': 8, 'sep': 9, 'sept': 9, 'septiembre': 9,
        'oct': 10, 'octubre': 10, 'nov': 11, 'noviembre': 11, 'dic': 12, 'diciembre': 12
    }
    mes = meses.get(mes_str)
    if mes is None:
        return None
    try:
        return datetime(anio, mes, dia)
    except:
        return None

def aplicar_formato_detalle(ws):
    fill_header = PatternFill(start_color="F4B183", end_color="F4B183", fill_type="solid")
    font_header = Font(name='Calibri', bold=True, size=11, color="000000")
    font_normal = Font(name='Calibri', size=10)
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    for cell in ws[1]:
        cell.fill = fill_header
        cell.font = font_header
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = thin_border
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=ws.max_column):
        for cell in row:
            cell.font = font_normal
            cell.border = thin_border
            cell.alignment = Alignment(vertical='center')
            if cell.column > 5:
                cell.number_format = '#,##0.00'
                cell.alignment = Alignment(horizontal='right', vertical='center')
            if cell.column == 1:
                cell.number_format = 'DD/MM/YYYY'
    anchos = {'A': 15, 'B': 18, 'C': 35, 'D': 18, 'E': 20, 'F': 25, 'G': 18, 'H': 18, 'I': 25, 'J': 22}
    for col, ancho in anchos.items():
        ws.column_dimensions[col].width = ancho
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions

def aplicar_formato_resumen(ws, fecha_inicio, fecha_fin):
    fill_header = PatternFill(start_color="F4B183", end_color="F4B183", fill_type="solid")
    fill_gral = PatternFill(start_color="92D050", end_color="92D050", fill_type="solid")
    fill_fecha = PatternFill(start_color="D9E2F3", end_color="D9E2F3", fill_type="solid")
    fill_celda = PatternFill(start_color="FFFFCC", end_color="FFFFCC", fill_type="solid")
    font_header = Font(name='Calibri', bold=True, size=11, color="000000")
    font_normal = Font(name='Calibri', size=10)
    font_total = Font(name='Calibri', bold=True, size=10)
    font_fecha = Font(name='Calibri', bold=True, size=11, color="000000")
    font_label = Font(name='Calibri', bold=True, size=10, color="000000")
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    
    ws.merge_cells('A1:H1')
    celda_titulo = ws['A1']
    celda_titulo.value = "RESUMEN POR TIENDA"
    celda_titulo.fill = fill_fecha
    celda_titulo.font = font_fecha
    celda_titulo.alignment = Alignment(horizontal='center', vertical='center')
    celda_titulo.border = thin_border
    for col in range(2, 9):
        ws.cell(row=1, column=col).fill = fill_fecha
        ws.cell(row=1, column=col).border = thin_border
    
    ws['A2'].value = "Fecha Inicio:"
    ws['A2'].font = font_label
    ws['A2'].alignment = Alignment(horizontal='right', vertical='center')
    ws['A2'].border = thin_border
    ws['B2'].value = fecha_inicio
    ws['B2'].font = font_normal
    ws['B2'].number_format = 'DD/MM/YYYY'
    ws['B2'].fill = fill_celda
    ws['B2'].alignment = Alignment(horizontal='center', vertical='center')
    ws['B2'].border = thin_border
    ws['C2'].value = "Fecha Fin:"
    ws['C2'].font = font_label
    ws['C2'].alignment = Alignment(horizontal='right', vertical='center')
    ws['C2'].border = thin_border
    ws['D2'].value = fecha_fin
    ws['D2'].font = font_normal
    ws['D2'].number_format = 'DD/MM/YYYY'
    ws['D2'].fill = fill_celda
    ws['D2'].alignment = Alignment(horizontal='center', vertical='center')
    ws['D2'].border = thin_border
    for col in range(5, 9):
        ws.cell(row=2, column=col).border = thin_border
    
    encabezados = ["Nombre de la tienda", "Suma de VENTAS", "Suma de Valor Ajustes Manuales",
                   "Suma de Comisión", "Suma de IVA Comisión (18%)", 
                   "Suma de Compensaciones", "Suma de Costo Canceladas", "Total"]
    for col_idx, encabezado in enumerate(encabezados, 1):
        cell = ws.cell(row=3, column=col_idx, value=encabezado)
        cell.fill = fill_header
        cell.font = font_header
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = thin_border
    
    tiendas = ws['A4']
    max_row = 4
    while tiendas.value is not None:
        tienda_nombre = tiendas.value
        if str(tienda_nombre).strip() == "TOTAL GENERAL":
            for col_idx in range(2, 9):
                col_letter = get_column_letter(col_idx)
                formula = f"=SUM({col_letter}4:{col_letter}{max_row-1})"
                cell = ws.cell(row=max_row, column=col_idx, value=formula)
                cell.font = font_total
                cell.fill = fill_gral
                cell.number_format = '#,##0.00'
                cell.alignment = Alignment(horizontal='right', vertical='center')
                cell.border = thin_border
            ws.cell(row=max_row, column=1).fill = fill_gral
            ws.cell(row=max_row, column=1).font = font_total
            ws.cell(row=max_row, column=1).border = thin_border
        else:
            ws.cell(row=max_row, column=2, value=f'=SUMIFS(Detalle!F:F,Detalle!A:A,">="&$B$2,Detalle!A:A,"<="&$D$2,Detalle!C:C,$A{max_row})')
            ws.cell(row=max_row, column=3, value=f'=SUMIFS(Detalle!J:J,Detalle!A:A,">="&$B$2,Detalle!A:A,"<="&$D$2,Detalle!C:C,$A{max_row})')
            ws.cell(row=max_row, column=4, value=f'=SUMIFS(Detalle!I:I,Detalle!A:A,">="&$B$2,Detalle!A:A,"<="&$D$2,Detalle!C:C,$A{max_row})')
            ws.cell(row=max_row, column=5, value=f'=D{max_row}*0.18')
            ws.cell(row=max_row, column=6, value=f'=SUMIFS(Detalle!G:G,Detalle!A:A,">="&$B$2,Detalle!A:A,"<="&$D$2,Detalle!C:C,$A{max_row})')
            ws.cell(row=max_row, column=7, value=f'=SUMIFS(Detalle!H:H,Detalle!A:A,">="&$B$2,Detalle!A:A,"<="&$D$2,Detalle!C:C,$A{max_row})')
            ws.cell(row=max_row, column=8, value=f'=SUM(B{max_row}:G{max_row})')
            
            for col_idx in range(2, 9):
                cell = ws.cell(row=max_row, column=col_idx)
                cell.font = font_normal
                cell.border = thin_border
                cell.alignment = Alignment(vertical='center')
                if col_idx > 1:
                    cell.number_format = '#,##0.00'
                    cell.alignment = Alignment(horizontal='right', vertical='center')
        
        ws.cell(row=max_row, column=1).font = font_normal
        ws.cell(row=max_row, column=1).border = thin_border
        ws.cell(row=max_row, column=1).alignment = Alignment(vertical='center')
        
        max_row += 1
        tiendas = ws.cell(row=max_row, column=1)
    
    anchos = {'A': 35, 'B': 18, 'C': 22, 'D': 18, 'E': 22, 'F': 22, 'G': 22, 'H': 18}
    for col, ancho in anchos.items():
        ws.column_dimensions[col].width = ancho
    ws.freeze_panes = 'A4'

def aplicar_formato_asiento(ws, fecha_inicio, fecha_fin):
    fill_header = PatternFill(start_color="F4B183", end_color="F4B183", fill_type="solid")
    fill_gral = PatternFill(start_color="92D050", end_color="92D050", fill_type="solid")
    fill_fecha = PatternFill(start_color="D9E2F3", end_color="D9E2F3", fill_type="solid")
    fill_celda = PatternFill(start_color="FFFFCC", end_color="FFFFCC", fill_type="solid")
    font_header = Font(name='Calibri', bold=True, size=11, color="000000")
    font_normal = Font(name='Calibri', size=10)
    font_total = Font(name='Calibri', bold=True, size=10)
    font_fecha = Font(name='Calibri', bold=True, size=11, color="000000")
    font_label = Font(name='Calibri', bold=True, size=10, color="000000")
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    
    ws.merge_cells('A1:G1')
    celda_titulo = ws['A1']
    celda_titulo.value = "ASIENTO CONTABLE"
    celda_titulo.fill = fill_fecha
    celda_titulo.font = font_fecha
    celda_titulo.alignment = Alignment(horizontal='center', vertical='center')
    celda_titulo.border = thin_border
    for col in range(2, 8):
        ws.cell(row=1, column=col).fill = fill_fecha
        ws.cell(row=1, column=col).border = thin_border
    
    ws['A2'].value = "Fecha Inicio:"
    ws['A2'].font = font_label
    ws['A2'].alignment = Alignment(horizontal='right', vertical='center')
    ws['A2'].border = thin_border
    ws['B2'].value = fecha_inicio
    ws['B2'].font = font_normal
    ws['B2'].number_format = 'DD/MM/YYYY'
    ws['B2'].fill = fill_celda
    ws['B2'].alignment = Alignment(horizontal='center', vertical='center')
    ws['B2'].border = thin_border
    ws['C2'].value = "Fecha Fin:"
    ws['C2'].font = font_label
    ws['C2'].alignment = Alignment(horizontal='right', vertical='center')
    ws['C2'].border = thin_border
    ws['D2'].value = fecha_fin
    ws['D2'].font = font_normal
    ws['D2'].number_format = 'DD/MM/YYYY'
    ws['D2'].fill = fill_celda
    ws['D2'].alignment = Alignment(horizontal='center', vertical='center')
    ws['D2'].border = thin_border
    for col in range(5, 8):
        ws.cell(row=2, column=col).border = thin_border
    
    encabezados = ["Tiendas", "Suma de Comision", 
                   "Suma de Descuento por inversión de Rappi a aplicar sobre Uso y alquiler de plataforma Rappi DAR",
                   "Comision_Abs_1", "Comision_Abs_2", "Code"]
    for col_idx, encabezado in enumerate(encabezados, 1):
        cell = ws.cell(row=3, column=col_idx, value=encabezado)
        cell.fill = fill_header
        cell.font = font_header
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = thin_border
    
    tiendas = ws['A4']
    max_row = 4
    while tiendas.value is not None:
        tienda_nombre = tiendas.value
        if str(tienda_nombre).strip() == "TOTAL GENERAL":
            for col_idx in [2, 4, 5]:
                col_letter = get_column_letter(col_idx)
                formula = f"=SUM({col_letter}4:{col_letter}{max_row-1})"
                cell = ws.cell(row=max_row, column=col_idx, value=formula)
                cell.font = font_total
                cell.fill = fill_gral
                cell.number_format = '#,##0.00'
                cell.alignment = Alignment(horizontal='right', vertical='center')
                cell.border = thin_border
            ws.cell(row=max_row, column=1).fill = fill_gral
            ws.cell(row=max_row, column=1).font = font_total
            ws.cell(row=max_row, column=1).border = thin_border
            ws.cell(row=max_row, column=3).fill = fill_gral
            ws.cell(row=max_row, column=3).border = thin_border
            ws.cell(row=max_row, column=6).fill = fill_gral
            ws.cell(row=max_row, column=6).border = thin_border
        else:
            ws.cell(row=max_row, column=2, value=f'=SUMIFS(Detalle!I:I,Detalle!A:A,">="&$B$2,Detalle!A:A,"<="&$D$2,Detalle!C:C,$A{max_row})')
            ws.cell(row=max_row, column=3, value="-")
            ws.cell(row=max_row, column=4, value=f'=ABS(B{max_row})')
            ws.cell(row=max_row, column=5, value=f'=ABS(B{max_row})')
            codigo = obtener_codigo_tienda(tienda_nombre)
            ws.cell(row=max_row, column=6, value=codigo)
            
            for col_idx in range(1, 7):
                cell = ws.cell(row=max_row, column=col_idx)
                cell.font = font_normal
                cell.border = thin_border
                cell.alignment = Alignment(vertical='center')
                if col_idx in [2, 4, 5]:
                    cell.number_format = '#,##0.00'
                    cell.alignment = Alignment(horizontal='right', vertical='center')
                elif col_idx in [3, 6]:
                    cell.alignment = Alignment(horizontal='center', vertical='center')
        
        ws.cell(row=max_row, column=1).font = font_normal
        ws.cell(row=max_row, column=1).border = thin_border
        ws.cell(row=max_row, column=1).alignment = Alignment(vertical='center')
        
        max_row += 1
        tiendas = ws.cell(row=max_row, column=1)
    
    anchos = {'A': 45, 'B': 18, 'C': 20, 'D': 18, 'E': 18, 'F': 15}
    for col, ancho in anchos.items():
        ws.column_dimensions[col].width = ancho
    ws.freeze_panes = 'A4'

if uploaded_file:
    if st.button("Procesar", use_container_width=True):
        with st.spinner("Procesando..."):
            try:
                df_completo = pd.read_excel(uploaded_file, sheet_name="Detalle", header=1)
                df_completo.columns = df_completo.columns.str.strip()
                
                columnas_necesarias = [
                    "Fecha de creación orden", "ID de la órden", "Nombre de la tienda",
                    "Estado de la órden", "Tipo de transacción",
                    "Ventas base por Uso y alquiler de plataforma Rappi (informativo)",
                    "Compensaciones", "Costo Canceladas",
                    "Uso y alquiler de plataforma Rappi", "Valor Ajustes Manuales"
                ]
                
                columnas_existentes = [col for col in columnas_necesarias if col in df_completo.columns]
                df_filtrado = df_completo[columnas_existentes].copy()
                
                if "Fecha de creación orden" in df_filtrado.columns:
                    df_filtrado["Fecha de creación orden"] = df_filtrado["Fecha de creación orden"].apply(convertir_fecha_manual)
                    df_filtrado = df_filtrado.dropna(subset=["Fecha de creación orden"])
                    
                    fecha_min = df_filtrado["Fecha de creación orden"].min()
                    fecha_max = df_filtrado["Fecha de creación orden"].max()
                    fecha_inicio_default = fecha_min
                    fecha_fin_default = fecha_max
                else:
                    fecha_inicio_default = datetime.now()
                    fecha_fin_default = datetime.now()
                
                if df_filtrado.empty:
                    st.warning("No hay registros validos.")
                    st.stop()
                
                tiendas_unicas = sorted(df_filtrado["Nombre de la tienda"].dropna().unique())
                
                resumen_data = [{"Nombre de la tienda": tienda} for tienda in tiendas_unicas]
                resumen_data.append({"Nombre de la tienda": "TOTAL GENERAL"})
                df_resumen = pd.DataFrame(resumen_data)
                
                asiento_data = [{"Tiendas": tienda} for tienda in tiendas_unicas]
                asiento_data.append({"Tiendas": "TOTAL GENERAL"})
                df_asiento = pd.DataFrame(asiento_data)
                
                excel_buffer = io.BytesIO()
                
                with pd.ExcelWriter(excel_buffer, engine='openpyxl') as writer:
                    df_filtrado.to_excel(writer, index=False, sheet_name='Detalle')
                    df_resumen.to_excel(writer, index=False, sheet_name='Resumen', startrow=3, header=False)
                    df_asiento.to_excel(writer, index=False, sheet_name='Asiento', startrow=3, header=False)
                    
                    aplicar_formato_detalle(writer.sheets['Detalle'])
                    aplicar_formato_resumen(writer.sheets['Resumen'], fecha_inicio_default, fecha_fin_default)
                    aplicar_formato_asiento(writer.sheets['Asiento'], fecha_inicio_default, fecha_fin_default)
                
                st.write(f"Registros procesados: {len(df_filtrado)}")
                
                st.download_button(
                    label="Guardar Excel",
                    data=excel_buffer.getvalue(),
                    file_name="rappi_asientos_contables.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    use_container_width=True
                )
                
            except Exception as e:
                st.error(f"Error: {e}")
