import streamlit as st
import pandas as pd
import io
import unicodedata
import re
from datetime import datetime
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

st.set_page_config(page_title="Filtrado Peya - Detalle y Resumen", layout="centered")
st.title("Filtrado de Datos Peya (Pedidos Ya)")

uploaded_files = st.file_uploader("Selecciona los reportes de Peya (puedes elegir varios)", type=['xlsx', 'xls'], accept_multiple_files=True)

# ==========================================
# MAPEO DE CÓDIGOS DE TIENDAS (PEYA)
# ==========================================
MAPEO_CODIGOS_PEYA = {
    "Fridays San Miguel": "0000000003",
    "Fridays Plaza Norte": "0000000017",
    "Fridays Mall Del Sur": "0000000012",
    "Fridays Trujillo": "0000000007",
    "Fridays Salaverry": "0000000008",
    "Fridays Óvalo Gutiérrez": "0000000001",
    "Fridays Primavera": "0000000009",
    "Fridays La Rambla San Borja": "0000000010",
    "Fridays Arequipa": "0000000006",
    "Fridays Santa Anita": "0000000015",
    "Fridays Dk La Molina": "0000000018",
    "Fridays Comas": "0000000016",
    "Fridays Puruchuco": "0000000014",
    "Fridays Mall Aventura San ... Lurigancho": "0000000020",
    "Smash Burger By Fridays Óvalo Monitor ,": "0000000019",
    "Smash Burger By Fridays Ma... Lurigancho": "0000000020",
    "Smash Burger By Fridays Óvalo Gutiérrez": "0000000001",
    "Fridays - Ovalo Monitor": "0000000019",
    "Smash Burger By Fridays Dk La Mol... ...": "0000000018",
    "Smash Burger By Fridays Santa Anita ,,,": "0000000015",
    "Smash Burger By Fridays - Santa Anita": "0000000015",
    "Smash Burger By Fridays - Dk La Molina": "0000000018",
    "Smash Burger By Fridays - Óv... Gutiérrez": "0000000001",
    "Smash Burger By Fridays - Óvalo Monitor": "0000000019",
    "Smash Burger By Fridays - ... Lurigancho": "0000000020",
    "Fridays La Rambla - B2c": "0000000010",
    "Fridays Mall Del Sur - B2c": "0000000012",
    "Fridays Óvalo Gutiérrez - B2c": "0000000001",
    "Fridays Real Plaza Salaverry - B2c": "0000000008",
    "Fridays - Real Plaza Piura": "0000000024",
    "Smash Burger By Fridays - Piura": "0000000024",
    "Fridays - Mall Aventura Porongoche": "0000000023",
    "Smash Burger By Fridays -... Porongoche.": "0000000023",
    "Smash Burger By Fridays - Re... Arequipa": "0000000006",
    "Smash Burger By Fridays - Rambla": "0000000010",
    "Smash Burger By Fridays - San Miguel": "0000000003",
    "Smash Burger By Fridays - Puruchuco": "0000000014",
    "Smash Burger By Fridays - Mall Del Sur": "0000000012",
    "Smash Burger By Fridays - Primavera": "0000000009",
    "Smash Burger By Fridays - Plaza Norte": "0000000017",
    "Smash Burger By Fridays - Salaverry": "0000000008",
    "Smash Burger By Fridays - Comas": "0000000016",
    "Fridays - El Polo": "0000000027",
    "Smash Burger By Fridays - El Polo": "0000000027",
    "Fridays Mall Aventura Santa Anita - B2c": "0000000015",
    "Smash Burger By Fridays - Ó... Gutiérrez": "0000000001",
    "Fridays Mall Plaza Comas - B2c": "0000000016",
    "Fridays El Polo - B2c": "0000000027",
    "Fridays - Cusco": "0000000025",
}

def normalizar_texto(texto):
    if pd.isna(texto):
        return ""
    texto = str(texto).strip()
    texto = unicodedata.normalize('NFKD', texto).encode('ascii', 'ignore').decode('utf-8')
    texto = re.sub(r'\s+', ' ', texto).lower()
    return texto

def obtener_codigo_tienda(nombre_tienda):
    if pd.isna(nombre_tienda):
        return "SIN_CODIGO"
    
    nombre_normalizado = normalizar_texto(nombre_tienda)
    
    for nombre_mapeado, codigo in MAPEO_CODIGOS_PEYA.items():
        if normalizar_texto(nombre_mapeado) == nombre_normalizado:
            return codigo
            
    for nombre_mapeado, codigo in MAPEO_CODIGOS_PEYA.items():
        nombre_mapeado_norm = normalizar_texto(nombre_mapeado)
        if nombre_normalizado in nombre_mapeado_norm or nombre_mapeado_norm in nombre_normalizado:
            return codigo
            
    return "SIN_CODIGO"

# ==========================================
# CONFIGURACIÓN DE COLUMNAS POR HOJA
# ==========================================
COLUMNAS_POR_HOJA = {
    "Lista de ordenes": ["Sucursal", "Fecha de pedido", "Monto comisionable Servicios PedidosYa", "Servicio Ventas PedidoYa ($)", "Cargos por Pedidos con Plus"],
    "Cargos por demoras": ["Sucursal", "Hora de retiro del pedido", "Cargo Total"],
    "Cargos por cancelaciones": ["Sucursal", "Fecha del pedido", "Monto del Pedido", "Monto final"],
    "Cargos por reclamos": ["Sucursal", "Motivo", "Monto del Pedido", "Cargos por reclamos de los usuarios"],
    "Reintegros": ["Sucursal", "Monto del Pedido", "% de Reintegro", "Monto neto a reintegrar"]
}

COLUMNAS_A_EXCLUIR_RESUMEN = {
    "Lista de ordenes": ["Fecha de pedido"],
    "Cargos por demoras": ["Hora de retiro del pedido"],
    "Cargos por cancelaciones": ["Fecha del pedido"],
    "Cargos por reclamos": ["Motivo"],
    "Reintegros": ["% de Reintegro"]
}

TABLAS_CON_FECHA = ["Lista de ordenes", "Cargos por cancelaciones"]

# ==========================================
# FUNCIÓN AUXILIAR PARA ESCRIBIR TABLAS EN DETALLE
# ==========================================
def agregar_tabla_detalle(writer, sheet_name, start_row, titulo, df):
    if df.empty:
        return start_row, None, None
    
    df.to_excel(writer, sheet_name=sheet_name, startrow=start_row + 1, index=False)
    worksheet = writer.sheets[sheet_name]
    
    title_font = Font(bold=True, size=14, color="4F81BD")
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center")
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

    cell_title = worksheet.cell(row=start_row + 1, column=1, value=titulo)
    cell_title.font = title_font
    
    header_row = start_row + 2
    last_col = len(df.columns)
    
    for col_idx in range(1, last_col + 1):
        cell = worksheet.cell(row=header_row, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border
        
    for r_idx in range(len(df)):
        for c_idx in range(1, last_col + 1):
            worksheet.cell(row=header_row + 1 + r_idx, column=c_idx).border = thin_border
            
    last_row_excel = header_row + len(df)
    last_col_letter = get_column_letter(last_col)
    worksheet.auto_filter.ref = f"A{header_row}:{last_col_letter}{last_row_excel}"
    
    for col_idx, col_name in enumerate(df.columns, start=1):
        max_length = len(str(col_name))
        for row in df[col_name]:
            cell_length = len(str(row))
            if cell_length > max_length:
                max_length = cell_length
        
        adjusted_width = min(max_length * 1.2 + 2, 50)
        col_letter = get_column_letter(col_idx)
        worksheet.column_dimensions[col_letter].width = adjusted_width
    
    data_start_row = header_row + 1
    data_end_row = last_row_excel
    
    return start_row + len(df) + 3, data_start_row, data_end_row

# ==========================================
# FUNCIÓN PARA CREAR FILA DE CONTROL ÚNICA EN RESUMEN
# ==========================================
def crear_fila_control(writer, sheet_name, fecha_inicio, fecha_fin):
    worksheet = writer.sheets[sheet_name]
    
    title_font = Font(bold=True, size=16, color="4F81BD")
    label_font = Font(bold=True, size=11)
    yellow_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
    
    cell_title = worksheet.cell(row=1, column=1, value="RESUMEN CONSOLIDADO")
    cell_title.font = title_font
    
    cell_label_inicio = worksheet.cell(row=2, column=1, value="Fecha Inicio:")
    cell_label_inicio.font = label_font
    cell_fecha_inicio = worksheet.cell(row=2, column=2)
    cell_fecha_inicio.value = fecha_inicio
    cell_fecha_inicio.fill = yellow_fill
    cell_fecha_inicio.number_format = 'DD/MM/YYYY'
    
    cell_label_fin = worksheet.cell(row=2, column=3, value="Fecha Fin:")
    cell_label_fin.font = label_font
    cell_fecha_fin = worksheet.cell(row=2, column=4)
    cell_fecha_fin.value = fecha_fin
    cell_fecha_fin.fill = yellow_fill
    cell_fecha_fin.number_format = 'DD/MM/YYYY'
    
    return "$B$2", "$D$2"

# ==========================================
# FUNCIÓN PARA ESCRIBIR TABLAS EN RESUMEN
# ==========================================
def agregar_tabla_resumen(writer, sheet_name, start_row, titulo, df_agrupado, tiene_fecha=False, 
                         detalle_start_row=None, detalle_end_row=None, col_fecha_detalle=None, 
                         col_montos_detalle=None, celda_fecha_inicio=None, celda_fecha_fin=None):
    if df_agrupado.empty:
        return start_row
    
    worksheet = writer.sheets[sheet_name]
    
    title_font = Font(bold=True, size=14, color="4F81BD")
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center")
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

    cell_title = worksheet.cell(row=start_row + 1, column=1, value=titulo)
    cell_title.font = title_font
    
    header_row = start_row + 2
    df_agrupado.to_excel(writer, sheet_name=sheet_name, startrow=header_row, index=False)
    
    last_col = len(df_agrupado.columns)
    
    for col_idx in range(1, last_col + 1):
        cell = worksheet.cell(row=header_row + 1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border
        
    num_sucursales = len(df_agrupado)
    for r_idx in range(num_sucursales):
        for c_idx in range(1, last_col + 1):
            cell = worksheet.cell(row=header_row + 2 + r_idx, column=c_idx)
            cell.border = thin_border
            
            if tiene_fecha and c_idx > 1 and c_idx < last_col:
                col_letra_detalle = get_column_letter(col_montos_detalle[c_idx - 2])
                col_fecha_letra = get_column_letter(col_fecha_detalle)
                
                rango_fecha = f"Detalle!${col_fecha_letra}${detalle_start_row}:${col_fecha_letra}${detalle_end_row}"
                rango_monto = f"Detalle!${col_letra_detalle}${detalle_start_row}:${col_letra_detalle}${detalle_end_row}"
                rango_sucursal = f"Detalle!$A${detalle_start_row}:$A${detalle_end_row}"
                celda_sucursal = f"A{header_row + 2 + r_idx}"
                
                formula = f'=SUMIFS({rango_monto},{rango_fecha},">="&{celda_fecha_inicio},{rango_fecha},"<="&{celda_fecha_fin},{rango_sucursal},{celda_sucursal})'
                cell.value = formula
            else:
                col_name = df_agrupado.columns[c_idx - 1]
                cell.value = df_agrupado.iloc[r_idx][col_name]
    
    for col_idx, col_name in enumerate(df_agrupado.columns, start=1):
        max_length = len(str(col_name))
        for row in df_agrupado[col_name]:
            cell_length = len(str(row))
            if cell_length > max_length:
                max_length = cell_length
        
        adjusted_width = min(max_length * 1.2 + 2, 50)
        col_letter = get_column_letter(col_idx)
        worksheet.column_dimensions[col_letter].width = adjusted_width
    
    return header_row + num_sucursales + 2

# ==========================================
# PROCESAMIENTO PRINCIPAL
# ==========================================
if uploaded_files:
    st.info(f"Archivos cargados: **{len(uploaded_files)}**")
    for f in uploaded_files:
        st.text(f"• {f.name}")
    
    if st.button("Generar Excel (Detalle y Resumen)", use_container_width=True, type="primary"):
        with st.spinner("Procesando y combinando todos los archivos..."):
            try:
                dfs_acumulados = {nombre_hoja: [] for nombre_hoja in COLUMNAS_POR_HOJA.keys()}
                
                for uploaded_file in uploaded_files:
                    excel_file = pd.ExcelFile(uploaded_file)
                    hojas_disponibles = excel_file.sheet_names
                    
                    for nombre_hoja, columnas_necesarias in COLUMNAS_POR_HOJA.items():
                        if nombre_hoja in hojas_disponibles:
                            df_hoja = pd.read_excel(excel_file, sheet_name=nombre_hoja)
                            columnas_existentes = [col for col in columnas_necesarias if col in df_hoja.columns]
                            if columnas_existentes:
                                dfs_acumulados[nombre_hoja].append(df_hoja[columnas_existentes])
                
                dfs_filtrados = {}
                fecha_min_global = None
                fecha_max_global = None
                
                for nombre_hoja, lista_dfs in dfs_acumulados.items():
                    if lista_dfs:
                        df_concat = pd.concat(lista_dfs, ignore_index=True)
                        
                        if nombre_hoja == "Lista de ordenes" and "Fecha de pedido" in df_concat.columns:
                            df_concat["Fecha de pedido"] = pd.to_datetime(df_concat["Fecha de pedido"], errors='coerce').dt.date
                            fechas_validas = df_concat["Fecha de pedido"].dropna()
                            if not fechas_validas.empty:
                                if fecha_min_global is None or fechas_validas.min() < fecha_min_global:
                                    fecha_min_global = fechas_validas.min()
                                if fecha_max_global is None or fechas_validas.max() > fecha_max_global:
                                    fecha_max_global = fechas_validas.max()
                        elif nombre_hoja == "Cargos por cancelaciones" and "Fecha del pedido" in df_concat.columns:
                            df_concat["Fecha del pedido"] = pd.to_datetime(df_concat["Fecha del pedido"], errors='coerce').dt.date
                            fechas_validas = df_concat["Fecha del pedido"].dropna()
                            if not fechas_validas.empty:
                                if fecha_min_global is None or fechas_validas.min() < fecha_min_global:
                                    fecha_min_global = fechas_validas.min()
                                if fecha_max_global is None or fechas_validas.max() > fecha_max_global:
                                    fecha_max_global = fechas_validas.max()
                        
                        dfs_filtrados[nombre_hoja] = df_concat
                
                if not dfs_filtrados:
                    st.error("No se pudo extraer información de ningún archivo.")
                    st.stop()
                
                # Si no hay fechas, usar fechas por defecto
                if fecha_min_global is None:
                    fecha_min_global = datetime(2020, 1, 1).date()
                if fecha_max_global is None:
                    fecha_max_global = datetime(2030, 12, 31).date()

                output = io.BytesIO()
                with pd.ExcelWriter(output, engine='openpyxl') as writer:
                    
                    start_row_detalle = 0
                    rangos_detalle = {}
                    
                    for nombre_hoja, df in dfs_filtrados.items():
                        tiene_fecha = nombre_hoja in TABLAS_CON_FECHA
                        nuevo_start, data_start, data_end = agregar_tabla_detalle(writer, "Detalle", start_row_detalle, nombre_hoja, df)
                        
                        if tiene_fecha and data_start is not None:
                            rangos_detalle[nombre_hoja] = {
                                'start_row': data_start,
                                'end_row': data_end,
                                'col_fecha': 2
                            }
                        
                        start_row_detalle = nuevo_start

                    writer.book.create_sheet("Resumen")
                    celda_fecha_inicio, celda_fecha_fin = crear_fila_control(writer, "Resumen", fecha_min_global, fecha_max_global)
                    
                    start_row_resumen = 3
                    sucursales_sin_codigo = set()

                    for nombre_hoja, df in dfs_filtrados.items():
                        df_resumen_data = df.copy()
                        columnas_excluir = COLUMNAS_A_EXCLUIR_RESUMEN.get(nombre_hoja, [])
                        cols_para_resumen = [col for col in df_resumen_data.columns if col not in columnas_excluir]
                        df_resumen_data = df_resumen_data[cols_para_resumen]
                        
                        for col in df_resumen_data.columns:
                            if col != 'Sucursal':
                                df_resumen_data[col] = pd.to_numeric(df_resumen_data[col], errors='coerce').fillna(0)
                        
                        df_agrupado = df_resumen_data.groupby('Sucursal', as_index=False).sum()
                        df_agrupado['Código Tienda'] = df_agrupado['Sucursal'].apply(obtener_codigo_tienda)
                        
                        sin_codigo_en_hoja = df_agrupado[df_agrupado['Código Tienda'] == 'SIN_CODIGO']['Sucursal'].unique()
                        sucursales_sin_codigo.update(sin_codigo_en_hoja)
                        
                        tiene_fecha = nombre_hoja in TABLAS_CON_FECHA
                        
                        detalle_start_row = None
                        detalle_end_row = None
                        col_fecha_detalle = None
                        col_montos_detalle = None
                        
                        if tiene_fecha and nombre_hoja in rangos_detalle:
                            detalle_start_row = rangos_detalle[nombre_hoja]['start_row']
                            detalle_end_row = rangos_detalle[nombre_hoja]['end_row']
                            col_fecha_detalle = rangos_detalle[nombre_hoja]['col_fecha']
                            col_montos_detalle = list(range(3, 3 + len([c for c in df_agrupado.columns if c not in ['Sucursal', 'Código Tienda']])))
                        
                        start_row_resumen = agregar_tabla_resumen(
                            writer, "Resumen", start_row_resumen, 
                            f"Resumen: {nombre_hoja}", df_agrupado,
                            tiene_fecha=tiene_fecha,
                            detalle_start_row=detalle_start_row,
                            detalle_end_row=detalle_end_row,
                            col_fecha_detalle=col_fecha_detalle,
                            col_montos_detalle=col_montos_detalle,
                            celda_fecha_inicio=celda_fecha_inicio,
                            celda_fecha_fin=celda_fecha_fin
                        )

                    if sucursales_sin_codigo:
                        st.warning(f"Se encontraron {len(sucursales_sin_codigo)} nombres de sucursal que no coinciden con el diccionario:")
                        for s in list(sucursales_sin_codigo)[:10]:
                            st.text(f"• '{s}'")
                        st.info("Si ves nombres aquí, cópialos y pégalos para que los agregue al diccionario.")
                    else:
                        st.success("¡Todas las sucursales fueron identificadas correctamente con su código!")

                output.seek(0)
                
                st.success(f"¡Excel generado exitosamente combinando {len(uploaded_files)} archivos!")
                st.download_button(
                    label="Descargar Excel Peya (Consolidado)",
                    data=output,
                    file_name="Peya_Detalle_y_Resumen_Consolidado.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    use_container_width=True
                )
                
            except Exception as e:
                st.error(f"Error durante el procesamiento: {e}")
                st.exception(e)
